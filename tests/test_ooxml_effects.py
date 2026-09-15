import json
import re
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from src.renderers.ppt_ooxml import (
    apply_gradient_fill,
    apply_gradient_line,
    apply_outer_shadow,
    apply_glow,
    apply_alpha,
    apply_style_effects,
    inspect_pptx_ooxml,
)
from src.style.planner import plan_style
from src.architecture.model import ArchitectureModel
from src.layout.planner import plan_layout
from src.renderers.ppt_writer import write_pptx


class OoxmlEffectsTests(unittest.TestCase):
    def test_gradient_shadow_glow_are_native_and_roundtrip(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "effects.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(1))
            apply_gradient_fill(shape, [{"position": 0, "color": "155EEF", "opacity": 1}, {"position": 1, "color": "7C3AED", "opacity": .75}])
            apply_gradient_line(shape, [{"position": 0, "color": "18B6D9"}, {"position": 1, "color": "7358E9"}], angle=90)
            apply_outer_shadow(shape, enabled=True)
            apply_glow(shape, enabled=True)
            prs.save(path)
            report = inspect_pptx_ooxml(path)
            self.assertTrue(report["pptx_package_valid"], report)
            self.assertTrue(report["slide_xml_valid"], report)
            self.assertGreaterEqual(report["gradient_fill_count"], 2, report)
            self.assertGreaterEqual(report["outer_shadow_count"], 1, report)
            self.assertGreaterEqual(report["glow_count"], 1, report)
            self.assertGreaterEqual(report["alpha_count"], 3, report)
            self.assertFalse(report["duplicate_fill_errors"], report)
            self.assertEqual(1, len(Presentation(path).slides))

    def test_line_alpha_and_gradient_alpha_are_native(self):
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
        shape.line.color.rgb = __import__("pptx").dml.color.RGBColor(0x12, 0x34, 0x56)
        apply_alpha(shape, 0.4, target="line")
        line = next(child for child in shape.element.spPr if child.tag.rsplit("}", 1)[-1] == "ln")
        alpha_values = [node.get("val") for node in line.iter() if node.tag.rsplit("}", 1)[-1] == "alpha"]
        self.assertEqual(["40000"], alpha_values)
        apply_gradient_line(shape, [{"position": 0, "color": "155EEF"}, {"position": 1, "color": "7C3AED"}])
        apply_alpha(shape, 0.5, target="line")
        alpha_values = [node.get("val") for node in line.iter() if node.tag.rsplit("}", 1)[-1] == "alpha"]
        self.assertEqual(["50000", "50000"], alpha_values)

    def test_disabled_effect_does_not_leave_empty_effect_list(self):
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
        apply_outer_shadow(shape, enabled=False)
        apply_glow(shape, enabled=False)
        self.assertFalse(any(child.tag.rsplit("}", 1)[-1] == "effectLst" for child in shape.element.spPr))

    def test_gradient_replacement_keeps_one_fill(self):
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
        shape.fill.solid(); shape.fill.fore_color.rgb = __import__("pptx").dml.color.RGBColor(255, 0, 0)
        apply_gradient_fill(shape, [{"position": 0, "color": "FFFFFF"}, {"position": 1, "color": "DDEBFF"}])
        fills = [c.tag.rsplit("}", 1)[-1] for c in shape.element.spPr if c.tag.rsplit("}", 1)[-1] in {"solidFill", "gradFill", "noFill"}]
        self.assertEqual(["gradFill"], fills)

    def test_invalid_gradient_falls_back_to_existing_solid_fill(self):
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
        shape.fill.solid(); shape.fill.fore_color.rgb = __import__("pptx").dml.color.RGBColor(255, 0, 0)
        apply_style_effects(shape, {"fill": {"type": "gradient", "stops": [{"position": 0, "color": "BAD"}]}})
        fills = [c.tag.rsplit("}", 1)[-1] for c in shape.element.spPr if c.tag.rsplit("}", 1)[-1] in {"solidFill", "gradFill", "noFill"}]
        self.assertEqual(["solidFill"], fills)

    def test_style_alias_and_tech_gradient_preset(self):
        style = plan_style("未来科技渐变风")
        self.assertEqual("tech_gradient", style["style_id"])
        self.assertEqual("SUPPORTED_OOXML", style["renderer_capabilities"]["gradient_fill"])
        self.assertEqual("gradient", style["header"]["fill"]["type"])
        self.assertEqual("gradient", style["node"]["fill"]["type"])

    def test_written_geometry_uses_integer_emu(self):
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "geometry.pptx"
            model = ArchitectureModel.from_dict({
                "title": "几何", "nodes": [{"id": "a", "title": "入口"}, {"id": "b", "title": "平台"}, {"id": "c", "title": "数据"}],
                "relations": [{"id": "r", "source": "a", "target": "c", "label": "汇聚"}],
            })
            write_pptx(model, plan_layout(model), plan_style("enterprise_tech"), out)
            with ZipFile(out) as package:
                xml = package.read("ppt/slides/slide1.xml").decode("utf-8")
            self.assertEqual([], re.findall(r'"-?\d+\.\d+"', xml), "OOXML numeric attributes must be integers")
            report = inspect_pptx_ooxml(out)
            self.assertEqual([], report["invalid_numeric_attributes"], report)
            # Reading geometry back raises when a float EMU slipped into the XML.
            for shape in Presentation(out).slides[0].shapes:
                self.assertIsInstance(int(shape.left + shape.top + shape.width + shape.height), int)

    def test_relation_label_sits_in_the_gap_between_shapes(self):
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "labelled.pptx"
            model = ArchitectureModel.from_dict({
                "title": "标签位置",
                "nodes": [
                    {"id": "a", "title": "服务对象", "layer": "服务对象层", "level": "L1", "priority": "primary"},
                    {"id": "b", "title": "平台能力", "layer": "平台支撑层", "level": "L3", "priority": "primary"},
                    {"id": "c", "title": "数据资源", "layer": "数据资源层", "level": "L4", "priority": "primary"},
                ],
                "relations": [{"id": "r", "source": "a", "target": "b", "label": "汇聚"}],
            })
            layout = plan_layout(model)
            relation = layout["relations"][0]
            self.assertGreaterEqual(len(relation["points"]), 2, relation)
            self.assertIn("label_box", relation)
            write_pptx(model, layout, plan_style("government_blue"), out)
            shapes = {s.text_frame.text: s for s in Presentation(out).slides[0].shapes if s.has_text_frame and s.text_frame.text}
            self.assertTrue({"服务对象", "平台能力", "数据资源", "汇聚"} <= set(shapes), shapes.keys())
            label, source, target = shapes["汇聚"], shapes["服务对象"], shapes["平台能力"]
            label_center = label.top + label.height / 2
            self.assertGreaterEqual(label_center, source.top + source.height, "label must not cover the source node")
            self.assertLessEqual(label_center, target.top, "label must not cover the target node")
            self.assertGreaterEqual(label.left, source.left + source.width / 2)

    def test_relations_carry_native_arrow_heads(self):
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "arrows.pptx"
            model = ArchitectureModel.from_dict({
                "title": "箭头",
                "nodes": [
                    {"id": "app", "title": "业务应用", "layer": "业务应用层", "level": "L2"},
                    {"id": "plat", "title": "平台能力", "layer": "平台支撑层", "level": "L3"},
                ],
                "relations": [{"id": "r", "source": "app", "target": "plat", "label": "调用"}],
            })
            layout = plan_layout(model)
            write_pptx(model, layout, plan_style("government_blue"), out)
            with ZipFile(out) as package:
                xml = package.read("ppt/slides/slide1.xml").decode("utf-8")
            self.assertEqual(1, len(re.findall(r"<a:tailEnd ", xml)), "each relation needs an arrow on the target end")
            self.assertGreaterEqual(len(re.findall(r"<p:cxnSp>", xml)), 1)

    def test_bands_and_layer_labels_are_native_shapes(self):
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "bands.pptx"
            model = ArchitectureModel.from_dict({
                "title": "层带",
                "nodes": [
                    {"id": "a", "title": "服务对象", "layer": "服务对象层", "level": "L1"},
                    {"id": "b", "title": "业务应用", "layer": "业务应用层", "level": "L2"},
                ],
                "relations": [],
            })
            layout = plan_layout(model)
            self.assertEqual(2, len(layout["bands"]))
            write_pptx(model, layout, plan_style("government_blue"), out)
            texts = {s.text_frame.text for s in Presentation(out).slides[0].shapes if s.has_text_frame}
            self.assertIn("服务对象层", texts, "band labels must be rendered as editable text")
            self.assertIn("业务应用层", texts)

    def test_primary_and_planned_nodes_differ_visually(self):
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "weight.pptx"
            model = ArchitectureModel.from_dict({
                "title": "权重",
                "nodes": [
                    {"id": "a", "title": "主节点", "layer": "平台支撑层", "level": "L3", "priority": "primary", "status": "current"},
                    {"id": "b", "title": "次节点", "layer": "平台支撑层", "level": "L3", "priority": "secondary", "status": "current"},
                    {"id": "c", "title": "待建节点", "layer": "平台支撑层", "level": "L3", "priority": "secondary", "status": "planned"},
                ],
                "relations": [],
            })
            write_pptx(model, plan_layout(model), plan_style("government_blue"), out)
            shapes = {s.text_frame.text: s for s in Presentation(out).slides[0].shapes if s.has_text_frame and s.text_frame.text}
            fills = {name: shapes[name].fill.fore_color.rgb for name in ("主节点", "次节点", "待建节点")}
            self.assertEqual(3, len(set(fills.values())), f"priority/status must change the fill: {fills}")
            self.assertEqual("DASH", str(shapes["待建节点"].line.dash_style).split(" ")[0], "planned nodes use a dashed border")

    def test_direct_writer_emits_editable_effects(self):
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "architecture.pptx"
            model = ArchitectureModel.from_dict({
                "title": "科技渐变架构", "nodes": [{"id": "a", "title": "入口"}, {"id": "b", "title": "平台"}],
                "relations": [{"id": "r", "source": "a", "target": "b", "label": "访问"}],
            })
            layout = plan_layout(model)
            write_pptx(model, layout, plan_style("tech_gradient"), out)
            report = inspect_pptx_ooxml(out)
            self.assertTrue(report["pptx_package_valid"], report)
            self.assertGreaterEqual(report["gradient_fill_count"], 2, report)
            self.assertGreaterEqual(report["outer_shadow_count"], 1, report)
            self.assertGreaterEqual(report["editable_shape_count"], 3, report)
            self.assertEqual(1, len(Presentation(out).slides))


if __name__ == "__main__":
    unittest.main()
