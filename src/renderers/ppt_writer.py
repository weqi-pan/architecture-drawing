"""Direct editable PowerPoint renderer with layer bands and an OOXML effects layer."""
from __future__ import annotations

from pathlib import Path
from typing import Any, Mapping

from src.architecture.model import ArchitectureModel
from src.renderers.ppt_ooxml import apply_style_effects, apply_arrow_head


class PPTDependencyError(RuntimeError):
    pass


def _mix(hex_a: str, hex_b: str, ratio: float) -> str:
    """Blend two six-digit hex colors; ratio is the weight of ``hex_b``."""
    def parts(value: str) -> tuple[int, int, int]:
        value = str(value or "FFFFFF").lstrip("#")
        value = "".join(ch * 2 for ch in value) if len(value) == 3 else (value + "FFFFFF")[:6]
        return tuple(int(value[i:i + 2], 16) for i in (0, 2, 4))
    a, b = parts(hex_a), parts(hex_b)
    return "".join(f"{round(a[i] + (b[i] - a[i]) * ratio):02X}" for i in range(3))


def write_pptx(model: ArchitectureModel | dict, layout: dict[str, Any], style: dict[str, Any], output: str | Path) -> Path:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise PPTDependencyError("Direct PPTX generation requires python-pptx; no Draw.io conversion is used.") from exc

    m = model if isinstance(model, ArchitectureModel) else ArchitectureModel.from_dict(model)
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    colors = style.get("colors", {})
    corner_radius = min(0.5, max(0.0, float(style.get("node_corner_radius", 0.08))))

    def rgb(value: str | None, fallback: str) -> RGBColor:
        value = (value or fallback).lstrip("#")
        return RGBColor.from_string((value + "FFFFFF")[:6])

    def coord(value: float) -> int:
        return Inches(float(value) / 96)

    def solid_color(fill: Any, fallback: str) -> str:
        if isinstance(fill, str):
            return fill
        if isinstance(fill, Mapping) and fill.get("type") == "solid":
            return str(fill.get("color", fallback))
        return fallback

    background = style.get("background", colors.get("background", "#FFFFFF"))
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(background, "FFFFFF")

    # Title band: the only element above the layer bands.
    header_style = dict(style.get("header", {}))
    if header_style:
        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, coord(38))
        base = solid_color(header_style.get("fill"), colors.get("primary", "26354A"))
        header.fill.solid(); header.fill.fore_color.rgb = rgb(base, "26354A")
        border = dict(header_style.get("border", {}))
        header.line.color.rgb = rgb(border.get("color"), colors.get("primary", "26354A")); header.line.width = Pt(float(border.get("width", 0.5)))
        tf = header.text_frame; tf.clear(); tf.margin_left = Inches(.22); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
        run = p.add_run(); run.text = m.title; run.font.size = Pt(18); run.font.bold = True; run.font.color.rgb = rgb("#FFFFFF", "FFFFFF")
        apply_style_effects(header, header_style, role="header")

    # Layer bands are native shapes so the layer structure stays visible and editable.
    band_style = dict(style.get("band", {}))
    band_fill = band_style.get("fill", _mix(colors.get("secondary", "E5ECF5"), "#FFFFFF", 0.35))
    band_border = band_style.get("border", _mix(colors.get("line", "718096"), "#FFFFFF", 0.55))
    band_label_color = band_style.get("label_color", colors.get("primary", "26354A"))
    for band in layout.get("bands", []):
        rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, coord(band["x"]), coord(band["y"]), coord(band["width"]), coord(band["height"]))
        rect.adjustments[0] = min(0.5, corner_radius + 0.04)
        rect.fill.solid(); rect.fill.fore_color.rgb = rgb(band_fill, "EFF3F8")
        rect.line.color.rgb = rgb(band_border, "D3DCE8"); rect.line.width = Pt(0.75)
        rect.shadow.inherit = False
        accent = band.get("accent", {})
        if accent:
            bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, coord(accent["x"]), coord(accent["y"]), coord(accent["width"]), coord(accent["height"]))
            bar.adjustments[0] = 0.5
            bar.fill.solid(); bar.fill.fore_color.rgb = rgb(band_label_color, "26354A")
            bar.line.fill.background(); bar.shadow.inherit = False
        label = band.get("label", {})
        if label:
            textbox = slide.shapes.add_textbox(coord(label["x"]), coord(label["y"]), coord(label["width"]), coord(label["height"]))
            frame = textbox.text_frame; frame.clear(); frame.word_wrap = True; frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            paragraph = frame.paragraphs[0]; paragraph.alignment = PP_ALIGN.LEFT
            run = paragraph.add_run(); run.text = str(band.get("name", ""))
            run.font.size = Pt(11); run.font.bold = True; run.font.color.rgb = rgb(band_label_color, "26354A")

    # Nodes: visual weight follows priority/status so layers are not a flat card wall.
    node_style = dict(style.get("node", {}))
    node_fill = solid_color(node_style.get("fill"), colors.get("secondary", "E5ECF5"))
    node_border = dict(node_style.get("border", {}))
    node_border_color = node_border.get("color", colors.get("primary", "26354A"))
    nodes = {n.id: n for n in m.nodes}; shapes = {}
    for item in layout.get("elements", []):
        n = nodes.get(item["id"])
        priority = str(getattr(n, "priority", "secondary") or "secondary").lower() if n else "secondary"
        status = str(getattr(n, "status", "current") or "current").lower() if n else "current"

        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, coord(item["x"]), coord(item["y"]), coord(item["width"]), coord(item["height"]))
        shape.adjustments[0] = corner_radius
        shape.shadow.inherit = False
        if priority == "primary":
            fill_color, border_color, text_color, bold = node_fill, node_border_color, colors.get("text_primary", node_border_color), True
        elif status in {"planned", "suggested"}:
            fill_color = _mix(node_fill, background, 0.55)
            border_color = colors.get("accent", node_border_color)
            text_color, bold = colors.get("text_secondary", node_border_color), False
        else:
            fill_color = _mix(node_fill, "#FFFFFF", 0.45)
            border_color = colors.get("line", node_border_color)
            text_color, bold = colors.get("text_secondary", node_border_color), False
        shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill_color, "E5ECF5")
        shape.line.color.rgb = rgb(border_color, "26354A")
        shape.line.width = Pt(float(node_border.get("width", 1.0)) * (1.35 if priority == "primary" else 0.85))
        if status in {"planned", "suggested"}:
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH

        tf = shape.text_frame; tf.clear(); tf.margin_left = Inches(.08); tf.margin_right = Inches(.08); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        run = p.add_run(); run.text = n.title if n else item["id"]
        run.font.size = Pt(14 if priority == "primary" else 12.5); run.font.bold = bold
        run.font.color.rgb = rgb(text_color, node_border_color)
        if n and n.subtitle:
            p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; p2.text = n.subtitle
            p2.font.size = Pt(9); p2.font.color.rgb = rgb(colors.get("text_secondary"), "64748B")
        apply_style_effects(shape, node_style, role="node")
        shapes[item["id"]] = shape

    # Relations: orthogonal routed segments with an arrow head pointing at the target.
    connector_style = dict(style.get("connector", {}))
    connector_color = connector_style.get("color", colors.get("line", "718096"))
    connector_width = float(connector_style.get("width", style.get("line_width", 1.25)))
    connector_effect = {"line": connector_style.get("line")} if connector_style.get("line") else {}
    for rel in layout.get("relations", []):
        a, b = shapes.get(rel["from"]), shapes.get(rel["to"])
        if a is None or b is None:
            continue
        points = rel.get("points") or []
        if len(points) < 2:
            continue
        segments = []
        for p1, p2 in zip(points, points[1:]):
            segments.append(slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, coord(p1[0]), coord(p1[1]), coord(p2[0]), coord(p2[1])))
        for index, line in enumerate(segments):
            line.line.color.rgb = rgb(connector_color, "718096")
            line.line.width = Pt(connector_width)
            if connector_effect:
                apply_style_effects(line, connector_effect, role="connector")
            if index == len(segments) - 1:
                apply_arrow_head(line, kind="tail")
        box = rel.get("label_box")
        if rel.get("label") and box:
            label = slide.shapes.add_textbox(coord(box["x"]), coord(box["y"]), coord(box["width"]), coord(box["height"]))
            # Solid background so the label reads as a break in the connector, not text on a line.
            label.fill.solid(); label.fill.fore_color.rgb = rgb(style.get("background", colors.get("background", "#FFFFFF")), "FFFFFF")
            label.line.fill.background(); label.shadow.inherit = False
            label.text_frame.word_wrap = False; label.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            label.text_frame.text = str(rel["label"])
            label.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            label.text_frame.paragraphs[0].font.size = Pt(8.5)
            label.text_frame.paragraphs[0].font.color.rgb = rgb(colors.get("text_secondary"), "64748B")

    out = Path(output); out.parent.mkdir(parents=True, exist_ok=True); prs.save(out); return out
